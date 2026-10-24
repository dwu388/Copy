from __future__ import annotations

from collections import defaultdict, deque
import math
from pathlib import Path
from typing import Literal

import pandas as pd
from openpyxl import load_workbook

EXCEL_COLUMNS = "K:P"
TRADE_HEADERS = [
    "selected_text",
    "action",
    "trader",
    "coin",
    "market_cap",
    "source_amount",
]

Order = Literal["top-to-bottom", "bottom-to-top"]
Pairing = Literal["fifo", "lifo"]


def _clean_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _number(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def load_trade_excel(path: Path, sheet_name: str | None = None, order: Order = "top-to-bottom") -> pd.DataFrame:
    """Load only the K:P trade contract from an .xlsx/.xlsm workbook.

    A:J and Q onward are intentionally ignored. K:P must contain the exact
    headers in TRADE_HEADERS so a shifted export fails loudly instead of
    silently training on the wrong columns.
    """
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError(f"Expected .xlsx or .xlsm input, got {path.suffix or '<no extension>'}")
    if order not in {"top-to-bottom", "bottom-to-top"}:
        raise ValueError(f"Unsupported order: {order}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is None:
            worksheet = workbook.worksheets[0]
        else:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"Worksheet {sheet_name!r} not found; available sheets: {workbook.sheetnames}"
                )
            worksheet = workbook[sheet_name]

        headers = [_clean_text(worksheet.cell(row=1, column=column).value) for column in range(11, 17)]
        if headers != TRADE_HEADERS:
            raise ValueError(
                "Excel K:P schema mismatch. Expected "
                f"{TRADE_HEADERS}, found {headers}. No fallback to A:D parsing is allowed."
            )

        records: list[dict[str, object]] = []
        for source_row, values in enumerate(
            worksheet.iter_rows(min_row=2, min_col=11, max_col=16, values_only=True), start=2
        ):
            if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
                continue
            selected_text, action, trader, coin, market_cap, source_amount = values
            records.append({
                "source_row": source_row,
                "selected_text": _clean_text(selected_text),
                "action": (_clean_text(action) or "").lower() or None,
                "trader": _clean_text(trader),
                "coin": _clean_text(coin),
                "market_cap": _number(market_cap),
                "source_amount": _number(source_amount),
            })
    finally:
        workbook.close()

    if order == "bottom-to-top":
        records.reverse()

    frame = pd.DataFrame(records, columns=["source_row", *TRADE_HEADERS])
    if not frame.empty:
        frame.insert(0, "sequence", range(1, len(frame) + 1))
    else:
        frame.insert(0, "sequence", pd.Series(dtype="int64"))
    return frame


def build_trade_cycles(
    frame: pd.DataFrame,
    *,
    trader: str | None = None,
    coin: str | None = None,
    pairing: Pairing = "fifo",
) -> pd.DataFrame:
    """Resolve each open BUY at the first later SELL for its trader and coin.

    The Android forecast labels price return from market cap, not the difference
    between two notification dollar amounts (which may describe partial fills).
    ``pairing`` is retained for CLI compatibility; first-sell resolution makes
    FIFO and LIFO equivalent because a sell closes all earlier open buys.
    """
    if pairing not in {"fifo", "lifo"}:
        raise ValueError(f"Unsupported pairing: {pairing}")

    required = {"sequence", "source_row", *TRADE_HEADERS}
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"Trade frame is missing columns: {sorted(missing)}")

    work = frame.copy()
    if trader is not None:
        wanted = trader.strip().casefold()
        work = work.loc[work["trader"].fillna("").astype(str).str.casefold().eq(wanted)]
    if coin is not None:
        wanted = coin.strip().casefold()
        work = work.loc[work["coin"].fillna("").astype(str).str.casefold().eq(wanted)]

    complete = (
        work["action"].isin(["bought", "sold"])
        & work["trader"].notna()
        & work["coin"].notna()
        & pd.to_numeric(work["market_cap"], errors="coerce").gt(0)
        & pd.to_numeric(work["source_amount"], errors="coerce").gt(0)
    )
    work = work.loc[complete].sort_values("sequence")

    open_buys: dict[tuple[str, str], deque] = defaultdict(deque)
    cycles: list[dict[str, object]] = []

    def close_cycle(buy: pd.Series, sell: pd.Series) -> dict[str, object]:
        buy_amount = float(buy["source_amount"])
        sell_amount = float(sell["source_amount"])
        roi = float(sell["market_cap"]) / float(buy["market_cap"]) - 1.0
        outcome = "win" if roi > 0 else ("loss" if roi < 0 else "breakeven")
        return {
            "trader": buy["trader"],
            "coin": buy["coin"],
            "status": "closed",
            "outcome": outcome,
            "is_win": int(roi > 0),
            "buy_sequence": int(buy["sequence"]),
            "sell_sequence": int(sell["sequence"]),
            "buy_row": int(buy["source_row"]),
            "sell_row": int(sell["source_row"]),
            "buy_selected_text": buy["selected_text"],
            "sell_selected_text": sell["selected_text"],
            "buy_market_cap": float(buy["market_cap"]),
            "sell_market_cap": float(sell["market_cap"]),
            "buy_source_amount": buy_amount,
            "sell_source_amount": sell_amount,
            # Dollar proceeds cannot be inferred from source notifications:
            # a SELL amount may be a partial sale or aggregate multiple buys.
            "profit_usd": pd.NA,
            "return_pct": roi * 100.0,
        }

    for _, row in work.iterrows():
        key = (str(row["trader"]).casefold(), str(row["coin"]).casefold())
        if row["action"] == "bought":
            open_buys[key].append(row)
            continue
        queue = open_buys[key]
        if not queue:
            continue
        while queue:
            buy = queue.popleft() if pairing == "fifo" else queue.pop()
            cycles.append(close_cycle(buy, row))

    for queue in open_buys.values():
        for buy in queue:
            cycles.append({
                "trader": buy["trader"],
                "coin": buy["coin"],
                "status": "open",
                "outcome": "open",
                "is_win": pd.NA,
                "buy_sequence": int(buy["sequence"]),
                "sell_sequence": pd.NA,
                "buy_row": int(buy["source_row"]),
                "sell_row": pd.NA,
                "buy_selected_text": buy["selected_text"],
                "sell_selected_text": None,
                "buy_market_cap": float(buy["market_cap"]),
                "sell_market_cap": pd.NA,
                "buy_source_amount": float(buy["source_amount"]),
                "sell_source_amount": pd.NA,
                "profit_usd": pd.NA,
                "return_pct": pd.NA,
            })

    columns = [
        "trader", "coin", "status", "outcome", "is_win", "buy_sequence", "sell_sequence",
        "buy_row", "sell_row", "buy_selected_text", "sell_selected_text",
        "buy_market_cap", "sell_market_cap", "buy_source_amount", "sell_source_amount",
        "profit_usd", "return_pct",
    ]
    result = pd.DataFrame(cycles, columns=columns)
    if not result.empty:
        result = result.sort_values("buy_sequence").reset_index(drop=True)
    return result
