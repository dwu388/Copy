from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from learning.trade_input import TRADE_HEADERS, build_trade_cycles, load_trade_excel


class TradeInputTest(unittest.TestCase):
    def _workbook(self, rows: list[list[object]], headers: list[str] | None = None) -> Path:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        path = Path(directory.name) / "trades.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Notifications"
        for column in range(1, 11):
            sheet.cell(1, column, f"ignored_{column}")
        for offset, header in enumerate(headers or TRADE_HEADERS, start=11):
            sheet.cell(1, offset, header)
        sheet.cell(1, 17, "ignored_17")
        for row_index, values in enumerate(rows, start=2):
            sheet.cell(row_index, 1, "must-not-be-read")
            for offset, value in enumerate(values, start=11):
                sheet.cell(row_index, offset, value)
            sheet.cell(row_index, 17, "must-not-be-read")
        workbook.save(path)
        workbook.close()
        return path

    def test_reader_uses_only_k_through_p_and_exact_headers(self) -> None:
        path = self._workbook([
            ["raw", "BOUGHT", " Aurelius0121 ", " BINF ", "1210000", "2034.62"],
        ])
        actual = load_trade_excel(path, sheet_name="Notifications")
        self.assertEqual(["sequence", "source_row", *TRADE_HEADERS], actual.columns.tolist())
        self.assertEqual("bought", actual.loc[0, "action"])
        self.assertEqual("Aurelius0121", actual.loc[0, "trader"])
        self.assertEqual("BINF", actual.loc[0, "coin"])
        self.assertEqual(1_210_000.0, actual.loc[0, "market_cap"])
        self.assertEqual(2034.62, actual.loc[0, "source_amount"])

    def test_shifted_or_old_schema_fails_instead_of_falling_back(self) -> None:
        headers = TRADE_HEADERS.copy()
        headers[0] = "text"
        path = self._workbook([], headers=headers)
        with self.assertRaisesRegex(ValueError, "K:P schema mismatch"):
            load_trade_excel(path)

    def test_first_sell_resolves_all_earlier_buys_by_price(self) -> None:
        path = self._workbook([
            ["orphan sell", "sold", "Aurelius0121", "BINF", 1_300_000, 2500.00],
            ["buy one", "bought", "Aurelius0121", "BINF", 1_210_000, 2034.62],
            ["buy two", "bought", "Aurelius0121", "BINF", 1_410_000, 2036.40],
            ["sell one", "sold", "Aurelius0121", "BINF", 1_600_000, 2400.00],
        ])
        trades = load_trade_excel(path)
        cycles = build_trade_cycles(trades, trader="aurelius0121", coin="binf")
        self.assertEqual(["closed", "closed"], cycles["status"].tolist())
        self.assertEqual("win", cycles.loc[0, "outcome"])
        self.assertEqual(1, cycles.loc[0, "is_win"])
        self.assertAlmostEqual((1_600_000 / 1_210_000 - 1) * 100, cycles.loc[0, "return_pct"])
        self.assertTrue(cycles["profit_usd"].isna().all())
        self.assertEqual(3, cycles.loc[0, "buy_row"])
        self.assertEqual(5, cycles.loc[0, "sell_row"])
        self.assertEqual(4, cycles.loc[1, "buy_row"])
        self.assertEqual("win", cycles.loc[1, "outcome"])
        self.assertEqual(5, cycles.loc[1, "sell_row"])

    def test_partial_sell_amount_cannot_reverse_price_outcome(self) -> None:
        path = self._workbook([
            ["buy", "bought", "trader", "COIN", 100_000, 1000],
            ["partial sell", "sold", "trader", "COIN", 120_000, 100],
        ])
        cycle = build_trade_cycles(load_trade_excel(path)).iloc[0]
        self.assertEqual(1, cycle["is_win"])
        self.assertAlmostEqual(20, cycle["return_pct"])

    def test_bottom_to_top_reverses_chronology(self) -> None:
        path = self._workbook([
            ["newest sell", "sold", "kch3n", "CRADLE", 1_570_000, 647.37],
            ["older buy", "bought", "kch3n", "CRADLE", 1_200_000, 498.32],
        ])
        trades = load_trade_excel(path, order="bottom-to-top")
        cycles = build_trade_cycles(trades)
        self.assertEqual(1, len(cycles))
        self.assertEqual("closed", cycles.loc[0, "status"])
        self.assertEqual(3, cycles.loc[0, "buy_row"])
        self.assertEqual(2, cycles.loc[0, "sell_row"])


if __name__ == "__main__":
    unittest.main()
